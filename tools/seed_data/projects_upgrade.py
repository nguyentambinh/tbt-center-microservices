"""
Họ tên: Nguyễn Tâm Bình
MSSV:   2324801030003
Môn:    Kiến trúc và thiết kế phần mềm
Bài:    Đồ án TBT Center đa quốc gia
Ngày:   2026-10-26

Nâng cấp dữ liệu dự án đang chạy: % khối lượng thực tế, nghiệm thu dự án đã quá hạn,
thêm dự án cấp 3, cấp 4 cho từng nước, tóm tắt và file Word cho mọi hợp đồng.

Đầu vào: thư mục JSON xuất từ CSDL từng khu vực (<kv>.pc.json, <kv>.hr.json, <kv>.acc.json).
Đầu ra: <out>/<kv>/project_contract.sql, finance.sql, files/contracts/<mã hợp đồng>.docx (tải lên tbt-media)

Cách dùng: python tools/seed_data/projects_upgrade.py <thư mục JSON> <thư mục ra> [ngày hôm nay YYYY-MM-DD]
"""
import datetime as dt
import json
import random
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

REGIONS = ['vn', 'us', 'de', 'br', 'za', 'au']
LANG = {'vn': 'vi', 'de': 'de', 'br': 'pt'}
RATE = {'VND': 25400, 'USD': 1, 'EUR': 0.92, 'BRL': 5.4, 'ZAR': 18.2, 'AUD': 1.52}
SALARY = {'VND': (2_000_000, 9_000_000), 'USD': (300, 1300), 'EUR': (280, 1200), 'BRL': (1500, 6500),
          'ZAR': (5000, 22000), 'AUD': (450, 1900)}
COMPANY = {'vn': 'TBT Center Việt Nam', 'us': 'TBT Center USA', 'de': 'TBT Center Deutschland GmbH',
           'br': 'TBT Center Brasil Ltda', 'za': 'TBT Center South Africa', 'au': 'TBT Center Australia'}

TEXT = {
    'vi': {
        'projects': ['Nhà máy điện mặt trời áp mái {c}', 'Trung tâm logistics {c}', 'Bệnh viện đa khoa {c}',
                     'Khu chung cư cao tầng {c}', 'Cầu vượt nút giao {c}', 'Nhà xưởng sản xuất linh kiện {c}'],
        'cities': ['Bình Dương', 'Long An', 'Đồng Nai', 'Bắc Ninh', 'Hải Phòng', 'Đà Nẵng'],
        'customers': ['Công ty CP Đầu tư Phú Hưng', 'Tập đoàn Hạ tầng Sao Mai', 'Công ty TNHH Logistics Đông Dương',
                      'Công ty CP Năng lượng Xanh Miền Nam', 'Bệnh viện Hoàn Mỹ Phương Nam', 'Công ty CP Địa ốc Thịnh Phát'],
        'types': {'1': 'Thi công trọn gói (EPC)', '2': 'Thi công phần thô', '3': 'Thiết kế và thi công', '4': 'Tư vấn giám sát'},
        'summary': 'Loại hợp đồng: {t}. Phạm vi: {name}. Thanh toán: tạm ứng {dep}%, các đợt theo nghiệm thu khối lượng, '
                   'giữ lại 5% bảo hành 12 tháng. Tiến độ: {start} đến {end}, nghiệm thu chậm nhất 60 ngày sau khi xong. '
                   'Phạt chậm tiến độ 0,05%/ngày, tối đa 8%. Bảo hành kết cấu 24 tháng.',
        'title': 'HỢP ĐỒNG THI CÔNG XÂY DỰNG', 'no': 'Số', 'party_a': 'Bên A (chủ đầu tư)', 'party_b': 'Bên B (nhà thầu)',
        'rows': ['Tên dự án', 'Loại hợp đồng', 'Giá trị hợp đồng', 'Tạm ứng', 'Ngày khởi công', 'Ngày hoàn thành', 'Nhân sự chủ chốt'],
        'articles': [
            'Điều 1. Phạm vi công việc: Bên B thực hiện {name} theo hồ sơ thiết kế được duyệt, bao gồm cung cấp vật tư, nhân công, máy thi công.',
            'Điều 2. Giá trị và thanh toán: tạm ứng {dep}% trong 10 ngày kể từ ngày ký; thanh toán theo từng đợt nghiệm thu khối lượng; giữ lại 5% giá trị làm bảo hành trong 12 tháng.',
            'Điều 3. Tiến độ: khởi công {start}, hoàn thành {end}. Chậm tiến độ do lỗi Bên B phạt 0,05% giá trị phần việc chậm cho mỗi ngày, tối đa 8% giá trị hợp đồng.',
            'Điều 4. Nghiệm thu: hai bên nghiệm thu từng giai đoạn và nghiệm thu bàn giao chậm nhất 60 ngày sau ngày hoàn thành.',
            'Điều 5. Bảo hành: 24 tháng phần kết cấu, 12 tháng phần hoàn thiện kể từ ngày nghiệm thu bàn giao.',
            'Điều 6. Điều khoản chung: tranh chấp giải quyết bằng thương lượng; không thành thì đưa ra Tòa án có thẩm quyền.'],
        'sign': ['ĐẠI DIỆN BÊN A', 'ĐẠI DIỆN BÊN B'],
    },
    'en': {
        'projects': ['Rooftop solar plant {c}', '{c} logistics hub', '{c} general hospital', '{c} residential tower',
                     '{c} interchange flyover', '{c} components factory'],
        'cities': ['Austin', 'Phoenix', 'Denver', 'Atlanta', 'Columbus', 'Charlotte'],
        'customers': ['Harbor Point Developments', 'Summit Infrastructure Group', 'BlueLine Logistics Inc.',
                      'GreenGrid Energy LLC', 'St. Mary Health Partners', 'Northgate Properties'],
        'types': {'1': 'Design-build (EPC)', '2': 'Structural works', '3': 'Design and construction', '4': 'Construction supervision'},
        'summary': 'Contract type: {t}. Scope: {name}. Payment: {dep}% advance, progress payments on certified works, '
                   '5% retention for 12 months. Schedule: {start} to {end}, handover acceptance within 60 days. '
                   'Delay penalty 0.05%/day up to 8%. Structural warranty 24 months.',
        'title': 'CONSTRUCTION CONTRACT', 'no': 'No.', 'party_a': 'Party A (owner)', 'party_b': 'Party B (contractor)',
        'rows': ['Project', 'Contract type', 'Contract value', 'Advance payment', 'Start date', 'Completion date', 'Key staff'],
        'articles': [
            'Article 1. Scope: Party B delivers {name} according to the approved design, including materials, labour and equipment.',
            'Article 2. Price and payment: {dep}% advance within 10 days of signing; progress payments on certified works; 5% retention for 12 months.',
            'Article 3. Schedule: start {start}, completion {end}. Delays caused by Party B incur 0.05% of the delayed works per day, capped at 8% of the contract value.',
            'Article 4. Acceptance: stage acceptance by both parties; final handover acceptance no later than 60 days after completion.',
            'Article 5. Warranty: 24 months for structural works and 12 months for finishing works from handover.',
            'Article 6. General: disputes are settled by negotiation, otherwise by the competent court.'],
        'sign': ['FOR PARTY A', 'FOR PARTY B'],
    },
}
TEXT['de'] = {**TEXT['en'],
              'projects': ['Solarkraftwerk auf Dach {c}', 'Logistikzentrum {c}', 'Klinikum {c}', 'Wohnturm {c}',
                           'Brückenkreuz {c}', 'Bauteilwerk {c}'],
              'cities': ['Leipzig', 'Dortmund', 'Nürnberg', 'Bremen', 'Hannover', 'Mainz'],
              'customers': ['Rheinwerk Immobilien AG', 'Nordbau Infrastruktur GmbH', 'Elbe Logistik GmbH',
                            'Sonnenfeld Energie GmbH', 'Klinikverbund Mitte', 'Stadtquartier Projekt GmbH'],
              'types': {'1': 'Generalunternehmer (EPC)', '2': 'Rohbau', '3': 'Planung und Bau', '4': 'Bauüberwachung'},
              'summary': 'Vertragsart: {t}. Umfang: {name}. Zahlung: {dep}% Anzahlung, Abschläge nach Abnahme, 5% Einbehalt '
                         '12 Monate. Termine: {start} bis {end}, Abnahme spätestens 60 Tage danach. Vertragsstrafe 0,05%/Tag, '
                         'max. 8%. Gewährleistung Tragwerk 24 Monate.',
              'title': 'BAUVERTRAG', 'no': 'Nr.', 'party_a': 'Auftraggeber', 'party_b': 'Auftragnehmer',
              'rows': ['Projekt', 'Vertragsart', 'Auftragswert', 'Anzahlung', 'Baubeginn', 'Fertigstellung', 'Schlüsselpersonal'],
              'sign': ['AUFTRAGGEBER', 'AUFTRAGNEHMER']}
TEXT['pt'] = {**TEXT['en'],
              'projects': ['Usina solar de telhado {c}', 'Centro logístico {c}', 'Hospital geral {c}', 'Torre residencial {c}',
                           'Viaduto {c}', 'Fábrica de componentes {c}'],
              'cities': ['Campinas', 'Curitiba', 'Recife', 'Goiânia', 'Joinville', 'Salvador'],
              'customers': ['Construtora Horizonte S.A.', 'Grupo Infra Sul', 'LogBrasil Armazéns Ltda',
                            'Energia Verde Nordeste', 'Hospital Santa Clara', 'Incorporadora Ipê'],
              'types': {'1': 'Empreitada global (EPC)', '2': 'Estrutura', '3': 'Projeto e obra', '4': 'Fiscalização'},
              'summary': 'Tipo: {t}. Escopo: {name}. Pagamento: {dep}% de adiantamento, medições por etapa, retenção de 5% '
                         'por 12 meses. Prazo: {start} a {end}, aceite final em até 60 dias. Multa 0,05%/dia, até 8%. '
                         'Garantia estrutural de 24 meses.',
              'title': 'CONTRATO DE OBRA', 'no': 'Nº', 'party_a': 'Contratante', 'party_b': 'Contratada',
              'rows': ['Projeto', 'Tipo de contrato', 'Valor', 'Adiantamento', 'Início', 'Conclusão', 'Equipe-chave'],
              'sign': ['CONTRATANTE', 'CONTRATADA']}
CITY_EN = {'za': ['Johannesburg', 'Durban', 'Pretoria', 'Port Elizabeth', 'Cape Town', 'Bloemfontein'],
           'au': ['Brisbane', 'Perth', 'Adelaide', 'Geelong', 'Newcastle', 'Hobart']}


def sql(v):
    if v is None:
        return 'null'
    if isinstance(v, (int, float)):
        return repr(v)
    if isinstance(v, (dt.date,)):
        return f"'{v.isoformat()}'"
    return "'" + str(v).replace("'", "''") + "'"


def day(s):
    return dt.date.fromisoformat(s) if s else None


def money_fmt(v, currency):
    return f'{v:,.0f} {currency}'.replace(',', '.' if currency in ('VND', 'EUR', 'BRL') else ',')


def write_contract(path, lang, c, kv):
    t = TEXT[lang]
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)
    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = head.add_run(t['title'])
    run.bold = True
    run.font.size = Pt(16)
    sub = doc.add_paragraph(f"{t['no']}: {c['id']}")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(f"{t['party_a']}: {c['customer']}")
    doc.add_paragraph(f"{t['party_b']}: {COMPANY[kv]}")
    table = doc.add_table(rows=0, cols=2)
    table.style = 'Table Grid'
    values = [c['name'], c['type_name'], money_fmt(c['value'], c['currency']) if c['value'] else '-',
              money_fmt(c['value'] * c['deposit'] / 100, c['currency']) if c['value'] else '-',
              c['start'].isoformat() if c['start'] else '-', c['end'].isoformat() if c['end'] else '-',
              ', '.join(c['staff']) or '-']
    for label, value in zip(t['rows'], values):
        cells = table.add_row().cells
        cells[0].text, cells[1].text = label, value
    doc.add_paragraph()
    for a in t['articles']:
        doc.add_paragraph(a.format(name=c['name'], dep=c['deposit'], start=c['start'], end=c['end']))
    sign = doc.add_table(rows=1, cols=2)
    for cell, text in zip(sign.rows[0].cells, t['sign']):
        cell.text = text
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)


def upgrade(kv, src, out, today, rnd):
    lang = LANG.get(kv, 'en')
    t = dict(TEXT[lang])
    if kv in CITY_EN:
        t['cities'] = CITY_EN[kv]
    data = json.loads((src / f'{kv}.pc.json').read_text(encoding='utf-8'))
    staff = json.loads((src / f'{kv}.hr.json').read_text(encoding='utf-8'))
    accepted_before = set(json.loads((src / f'{kv}.acc.json').read_text(encoding='utf-8')))
    projects = data['projects']
    currency = projects[0]['currency']
    country = projects[0]['country']
    rate = RATE[currency]
    images = [p['image'] for p in projects if p['image']]
    companies = sorted({p['company'] for p in projects})
    pc, fin = [], []
    contracts_out = []
    members_of = {}

    # Dự án cũ: % khối lượng thực tế; quá hạn thì phần lớn đã nghiệm thu
    for p in projects:
        if p['status'] == 'accepted':
            continue
        start, end = day(p['start']), day(p['end'])
        if start and start > today:
            percent = 0
        elif end and end < today:
            if rnd.random() < 0.6:
                accepted_on = min(today, end + dt.timedelta(days=rnd.randint(5, 50)))
                pc.append(f"update projects set status = 'accepted', accepted_on = {sql(accepted_on)}, "
                          f"accepted_by = 'finance', progress_percent = 100 where id = {sql(p['id'])};")
                if p['id'] not in accepted_before and p['value']:
                    fin.append(acceptance_sql(p['id'], accepted_on, float(p['value']), currency, rate))
                continue
            percent = rnd.randint(78, 96)
        elif start and end:
            frac = (today - start).days / max(1, (end - start).days)
            percent = max(3, min(97, round(frac * 100 + rnd.randint(-15, 8))))
        else:
            percent = rnd.randint(10, 60)
        pc.append(f'update projects set progress_percent = {percent} where id = {sql(p["id"])};')

    # Dự án mới cấp 3 (300k–800k USD) và cấp 4 (từ 800k USD), đủ trạng thái
    plans = [  # (cấp, bắt đầu cách hôm nay, số ngày làm, trạng thái)
        (3, -420, 360, 'accepted_recent'), (4, -700, 640, 'accepted_last_month'), (3, -200, 420, 'in_progress'),
        (4, -90, 720, 'in_progress'), (3, -330, 300, 'overdue'), (4, 45, 540, 'not_started'),
    ]
    customer_rows, contract_rows, project_rows, link_rows, member_rows = [], [], [], [], []
    for i, (tier, offset, length, state) in enumerate(plans, 1):
        usd = rnd.randint(320_000, 760_000) if tier == 3 else rnd.randint(900_000, 2_600_000)
        value = round(usd * rate, -3 if rate > 100 else 0)
        start = today + dt.timedelta(days=offset)
        end = start + dt.timedelta(days=length)
        code = f'{kv.upper()}X{i:03d}'
        customer_id, contract_id, link_id = f'{code}KH', f'{code}HD', f'{code}PC'
        name = t['projects'][(i - 1) % 6].format(c=t['cities'][(i - 1) % 6])
        customer = t['customers'][(i - 1) % 6]
        ctype = str(rnd.randint(1, 3))
        deposit = rnd.choice([10, 15, 20, 30])
        status, accepted_on, percent = 'in_progress', None, None
        if state == 'accepted_recent':
            status, accepted_on, percent = 'accepted', today - dt.timedelta(days=rnd.randint(3, 20)), 100
        elif state == 'accepted_last_month':
            status, accepted_on, percent = 'accepted', (today.replace(day=1) - dt.timedelta(days=rnd.randint(3, 25))), 100
        elif state == 'not_started':
            percent = 0
        elif state == 'overdue':
            percent = rnd.randint(82, 94)
        else:
            frac = (today - start).days / length
            percent = max(5, min(95, round(frac * 100 + rnd.randint(-8, 8))))
        customer_rows.append(f"({sql(customer_id)}, {sql(country)}, 'business', {sql(customer)}, null, null)")
        contract_rows.append(f"({sql(contract_id)}, {sql(country)}, {sql(name)}, {sql(ctype)}, {sql(start)}, {sql(end)})")
        project_rows.append(f"({sql(code)}, {sql(rnd.choice(companies))}, {sql(country)}, {sql(customer_id)}, "
                            f"{sql(round(value * deposit / 100, 2))}, {sql(float(value))}, {sql(currency)}, {sql(status)}, "
                            f"{sql(accepted_on)}, {sql('finance' if accepted_on else None)}, "
                            f"{sql(rnd.choice(images) if images else None)}, {sql(percent)})")
        link_rows.append(f"({sql(link_id)}, {sql(code)}, {sql(contract_id)}, {sql(end)})")
        team = rnd.sample(staff, min(len(staff), rnd.randint(5, 9)))
        members_of[contract_id] = [m['name'] for m in team[:3]]
        lo, hi = SALARY[currency]
        for m in team:
            hours = 0 if state == 'not_started' else rnd.randint(60, 420)
            member_rows.append(f"({sql(link_id)}, {sql(m['id'])}, {hours}, {sql(float(rnd.randint(lo, hi)))})")
        if accepted_on:
            fin.append(acceptance_sql(code, accepted_on, float(value), currency, rate))
        contracts_out.append({'id': contract_id, 'name': name, 'type': ctype, 'start': start, 'end': end,
                              'customer': customer, 'value': float(value), 'deposit': deposit, 'currency': currency})

    pc.append('insert into customers (id, country_code, customer_type, name, contact, address) values\n  '
              + ',\n  '.join(customer_rows) + '\non conflict (id) do nothing;')
    pc.append('insert into contracts (id, country_code, name, contract_type, start_date, expected_end_date) values\n  '
              + ',\n  '.join(contract_rows) + '\non conflict (id) do nothing;')
    pc.append('insert into projects (id, company_id, country_code, customer_id, deposit_amount, expected_acceptance_amount, '
              'currency, status, accepted_on, accepted_by, image_key, progress_percent) values\n  '
              + ',\n  '.join(project_rows) + '\non conflict (id) do nothing;')
    pc.append('insert into project_contracts (id, project_id, contract_id, end_date) values\n  '
              + ',\n  '.join(link_rows) + '\non conflict (id) do nothing;')
    pc.append('insert into project_members (project_contract_id, employee_id, total_hours, project_salary) values\n  '
              + ',\n  '.join(member_rows) + '\non conflict do nothing;')

    # Hợp đồng cũ: lấy giá trị, khách hàng của dự án gắn với hợp đồng
    by_project = {p['id']: p for p in projects}
    customers = {c['id']: c['name'] for c in data['customers']}
    for c in data['contracts']:
        p = by_project.get(c['project'], {})
        contracts_out.append({'id': c['id'], 'name': c['name'], 'type': c['type'] or '1', 'start': day(c['start']),
                              'end': day(c['end']), 'customer': customers.get(p.get('customer'), '-'),
                              'value': float(p['value']) if p.get('value') else 0, 'deposit': 30, 'currency': currency})

    for c in contracts_out:
        c['type_name'] = t['types'].get(str(c['type']), t['types']['1'])
        c['staff'] = members_of.get(c['id'], [])
        summary = t['summary'].format(t=c['type_name'], name=c['name'], dep=c['deposit'], start=c['start'], end=c['end'])
        key = f"contracts/{c['id']}.docx"
        # Hợp đồng đã có file thật thì giữ file đó
        pc.append(f"update contracts set summary = {sql(summary)}, file_key = coalesce(file_key, {sql(key)}) "
                  f"where id = {sql(c['id'])};")
        write_contract(out / kv / 'files' / key, lang, c, kv)

    region_out = out / kv
    region_out.mkdir(parents=True, exist_ok=True)
    (region_out / 'project_contract.sql').write_text('begin;\n' + '\n'.join(pc) + '\ncommit;\n', encoding='utf-8')
    (region_out / 'finance.sql').write_text('begin;\n' + '\n'.join(fin) + '\ncommit;\n', encoding='utf-8')
    return len(plans), len(contracts_out), len(fin)


def acceptance_sql(project_id, on, amount, currency, rate):
    return ('insert into acceptance_revenues (project_id, accepted_on, amount_original, currency_original, rate_applied, '
            f'rate_date, amount_base) values ({sql(project_id)}, {sql(on)}, {amount:.2f}, {sql(currency)}, {rate}, '
            f'{sql(on)}, {amount / rate:.2f}) on conflict (project_id) do nothing;')


def main():
    src, out = Path(sys.argv[1]), Path(sys.argv[2])
    today = dt.date.fromisoformat(sys.argv[3]) if len(sys.argv) > 3 else dt.date.today()
    for kv in REGIONS:
        added, contracts, revenues = upgrade(kv, src, out, today, random.Random(f'tbt-{kv}-projects'))
        print(f'{kv}: thêm {added} dự án, {contracts} hợp đồng có file, {revenues} khoản doanh thu nghiệm thu')


if __name__ == '__main__':
    main()
